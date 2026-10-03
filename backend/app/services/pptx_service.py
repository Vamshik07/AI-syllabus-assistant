import os
import logging
from typing import Dict, List, Any
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from app.config import settings

logger = logging.getLogger(__name__)

class PPTXService:
    def extract_document(self, file_path: str, document_id: str) -> Dict[str, Any]:
        """
        Parses a PPTX file using python-pptx.
        Preserves slide order, slide titles, text boxes, images, tables, shape positions,
        and layout information.
        """
        prs = Presentation(file_path)
        slide_count = len(prs.slides)
        slides_data = []
        full_text = []

        doc_media_dir = os.path.join(settings.EXTRACTED_MEDIA_DIR, document_id)
        os.makedirs(doc_media_dir, exist_ok=True)

        for slide_idx, slide in enumerate(prs.slides):
            slide_number = slide_idx + 1
            slide_title = ""
            if slide.shapes.title and slide.shapes.title.text:
                slide_title = slide.shapes.title.text.strip()

            elements = []
            slide_text_pieces = []
            if slide_title:
                slide_text_pieces.append(f"Title: {slide_title}")

            for shape_idx, shape in enumerate(slide.shapes):
                left = getattr(shape, "left", 0)
                top = getattr(shape, "top", 0)
                width = getattr(shape, "width", 0)
                height = getattr(shape, "height", 0)

                # Skip title shape if already captured as title
                if shape == slide.shapes.title:
                    continue

                # 1. Text Frame
                if shape.has_text_frame:
                    paragraphs = []
                    for p in shape.text_frame.paragraphs:
                        p_text = p.text.strip()
                        if p_text:
                            is_bold = any(run.font.bold for run in p.runs if run.font.bold is not None)
                            paragraphs.append({
                                "text": p_text,
                                "level": p.level,
                                "is_bold": is_bold,
                            })
                            slide_text_pieces.append(p_text)

                    if paragraphs:
                        combined_text = "\n".join(p["text"] for p in paragraphs)
                        elements.append({
                            "type": "TEXT",
                            "text": combined_text,
                            "paragraphs": paragraphs,
                            "x1": float(left),
                            "y1": float(top),
                            "x2": float(left + width),
                            "y2": float(top + height),
                        })

                # 2. Table
                elif shape.has_table:
                    table = shape.table
                    rows_data = []
                    for row in table.rows:
                        row_cells = [cell.text.strip() for cell in row.cells]
                        rows_data.append(row_cells)
                    
                    html_table = self._convert_table_to_html(rows_data)
                    table_text = " | ".join([" ".join(r) for r in rows_data])
                    slide_text_pieces.append(table_text)
                    elements.append({
                        "type": "TABLE",
                        "text": table_text,
                        "html": html_table,
                        "x1": float(left),
                        "y1": float(top),
                        "x2": float(left + width),
                        "y2": float(top + height),
                    })

                # 3. Image
                elif shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                    try:
                        image = shape.image
                        img_filename = f"slide_{slide_number}_img_{shape_idx}.{image.ext}"
                        img_path = os.path.join(doc_media_dir, img_filename)
                        with open(img_path, "wb") as f:
                            f.write(image.blob)
                        
                        rel_img_path = f"/api/media/{document_id}/{img_filename}"
                        
                        # Extract diagram / educational content from the image
                        from app.services.ai_service import ai_service
                        img_data = ai_service.extract_image_content(img_path)
                        img_heading = img_data.get("heading")
                        img_content = img_data.get("content", "").strip()

                        if not slide_title and img_heading:
                            slide_title = img_heading

                        img_text = img_content if img_content else f"[Image on Slide {slide_number}]"
                        if img_content:
                            slide_text_pieces.append(f"Diagram Content ({img_heading or 'Diagram'}):\n{img_content}")

                        elements.append({
                            "type": "IMAGE",
                            "media_path": rel_img_path,
                            "text": img_text,
                            "diagram_heading": img_heading,
                            "x1": float(left),
                            "y1": float(top),
                            "x2": float(left + width),
                            "y2": float(top + height),
                        })
                    except Exception as e:
                        logger.warning(f"Failed to extract PPTX image: {e}")

            if not slide_title:
                # Check for first heading / bold paragraph
                for el in elements:
                    if el.get("type") == "TEXT" and el.get("paragraphs"):
                        candidate = el["paragraphs"][0]["text"].strip()
                        if candidate and len(candidate) <= 60 and not candidate.startswith("---"):
                            slide_title = candidate
                            break

            raw_slide_text = "\n".join(slide_text_pieces).strip()
            full_text.append(f"--- Slide {slide_number}: {slide_title} ---\n{raw_slide_text}")

            slides_data.append({
                "slide_number": slide_number,
                "title": slide_title,
                "elements": elements,
                "text": raw_slide_text,
            })

        return {
            "slide_count": slide_count,
            "slides": slides_data,
            "full_text": "\n\n".join(full_text),
        }

    def _convert_table_to_html(self, rows: List[List[str]]) -> str:
        if not rows:
            return ""
        html = ["<table class='min-w-full border border-gray-300 divide-y divide-gray-200 text-sm'>"]
        header = rows[0]
        html.append("<thead class='bg-gray-50'><tr>")
        for cell in header:
            html.append(f"<th class='px-3 py-2 text-left font-semibold text-gray-700 border-b'>{cell}</th>")
        html.append("</tr></thead>")
        html.append("<tbody class='divide-y divide-gray-200 bg-white'>")
        for row in rows[1:]:
            html.append("<tr>")
            for cell in row:
                html.append(f"<td class='px-3 py-2 border-b text-gray-800'>{cell}</td>")
            html.append("</tr>")
        html.append("</tbody></table>")
        return "".join(html)

pptx_service = PPTXService()
