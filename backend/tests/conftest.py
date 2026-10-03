import os
import io
import pytest
from fastapi.testclient import TestClient
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool
import docx
from pptx import Presentation
from pptx.util import Inches, Pt
import pymupdf as fitz

from app.database.database import Base, get_db
from app.main import app
from app.models.models import User, Subject, Unit

# Use in-memory SQLite for fast, isolated test suite
SQLALCHEMY_DATABASE_URL = "sqlite:///:memory:"

engine = create_engine(
    SQLALCHEMY_DATABASE_URL,
    connect_args={"check_same_thread": False},
    poolclass=StaticPool,
)
TestingSessionLocal = sessionmaker(autocommit=False, autoflush=False, bind=engine)

@pytest.fixture(scope="function")
def db_session():
    Base.metadata.create_all(bind=engine)
    session = TestingSessionLocal()
    # Seed default user
    user = User(id="test-user-id", email="student@sourcelearn.ai", name="Test Student")
    session.add(user)
    session.commit()
    try:
        yield session
    finally:
        session.close()
        Base.metadata.drop_all(bind=engine)

@pytest.fixture(scope="function")
def client(db_session):
    def override_get_db():
        try:
            yield db_session
        finally:
            pass

    app.dependency_overrides[get_db] = override_get_db
    with TestClient(app) as test_client:
        yield test_client
    app.dependency_overrides.clear()

@pytest.fixture
def sample_pdf_bytes():
    """Generates a real valid multi-page PDF document in memory."""
    doc = fitz.open()
    # Page 1: Chapter 1 Introduction
    p1 = doc.new_page()
    p1.insert_text((50, 72), "1. Introduction to Networks", fontsize=18, fontname="helv")
    p1.insert_text((50, 110), "Computer networks allow computers to share resources.", fontsize=11, fontname="helv")
    p1.insert_text((50, 130), "1.1 Computer Basics", fontsize=14, fontname="helv")
    p1.insert_text((50, 150), "A computer system contains hardware and software components.", fontsize=11, fontname="helv")
    
    # Page 2: Chapter 2 Network Models
    p2 = doc.new_page()
    p2.insert_text((50, 72), "2. Network Models", fontsize=18, fontname="helv")
    p2.insert_text((50, 110), "The OSI model specifies seven architectural communication layers.", fontsize=11, fontname="helv")
    p2.insert_text((50, 130), "2.1 TCP/IP Model", fontsize=14, fontname="helv")
    p2.insert_text((50, 150), "The TCP/IP model has four operational layers.", fontsize=11, fontname="helv")

    buf = io.BytesIO()
    doc.save(buf)
    doc.close()
    return buf.getvalue()

@pytest.fixture
def sample_pptx_bytes():
    """Generates a real valid PPTX document in memory."""
    prs = Presentation()
    # Slide 1
    s1_layout = prs.slide_layouts[0]
    slide1 = prs.slides.add_slide(s1_layout)
    slide1.shapes.title.text = "Introduction to Computing"
    subtitle = slide1.placeholders[1]
    subtitle.text = "Computing involves systematic data processing and automated calculations."

    # Slide 2
    s2_layout = prs.slide_layouts[1]
    slide2 = prs.slides.add_slide(s2_layout)
    slide2.shapes.title.text = "Computer Architecture"
    body = slide2.placeholders[1]
    tf = body.text_frame
    tf.text = "Von Neumann architecture separates memory and arithmetic processing."

    # Slide 3
    s3_layout = prs.slide_layouts[1]
    slide3 = prs.slides.add_slide(s3_layout)
    slide3.shapes.title.text = "Operating Systems"
    body3 = slide3.placeholders[1]
    body3.text_frame.text = "Operating systems manage computer hardware and execute user applications."

    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()

@pytest.fixture
def sample_docx_bytes():
    """Generates a real valid DOCX document with native headings in memory."""
    doc = docx.Document()
    # Topic 1
    doc.add_heading("1. Operating System Fundamentals", level=1)
    doc.add_paragraph("An operating system acts as an intermediary between a user and computer hardware.")
    doc.add_heading("Process Management", level=2)
    doc.add_paragraph("A process is a program in execution requiring CPU and memory resources.")

    # Topic 2
    doc.add_heading("2. Memory Management", level=1)
    doc.add_paragraph("Memory management coordinates physical memory and virtual address spaces.")
    doc.add_heading("Virtual Memory", level=2)
    doc.add_paragraph("Virtual memory uses paging to map virtual addresses to physical frames.")

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
