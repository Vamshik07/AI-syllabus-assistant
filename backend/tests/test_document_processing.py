import pytest
import io

def test_invalid_file_uploads(client):
    # Setup subject and unit
    s_res = client.post("/api/subjects", json={"name": "Test Subject"})
    subj_id = s_res.json()["id"]
    u_res = client.post(f"/api/subjects/{subj_id}/units", json={"title": "Test Unit"})
    unit_id = u_res.json()["id"]

    # 1. Invalid extension .txt
    res = client.post(
        f"/api/units/{unit_id}/document",
        files={"file": ("test.txt", b"plain text content", "text/plain")}
    )
    assert res.status_code == 400
    assert "Unsupported file extension" in res.json()["detail"]

    # 2. Invalid extension .exe
    res = client.post(
        f"/api/units/{unit_id}/document",
        files={"file": ("malicious.exe", b"MZ\x90\x00\x03", "application/x-msdownload")}
    )
    assert res.status_code == 400

    # 3. Empty file
    res = client.post(
        f"/api/units/{unit_id}/document",
        files={"file": ("empty.pdf", b"", "application/pdf")}
    )
    assert res.status_code == 400
    assert "empty" in res.json()["detail"].lower()

    # 4. Corrupted PDF (missing %PDF- header)
    res = client.post(
        f"/api/units/{unit_id}/document",
        files={"file": ("corrupted.pdf", b"NOT_A_REAL_PDF_HEADER_DATA", "application/pdf")}
    )
    assert res.status_code == 400
    assert "corrupted" in res.json()["detail"].lower()

def test_pdf_processing(client, sample_pdf_bytes):
    s_res = client.post("/api/subjects", json={"name": "Computer Networks"})
    u_res = client.post(f"/api/subjects/{s_res.json()['id']}/units", json={"title": "Unit 1: Networks Overview"})
    unit_id = u_res.json()["id"]

    # Upload PDF
    upload_res = client.post(
        f"/api/units/{unit_id}/document",
        files={"file": ("networks.pdf", sample_pdf_bytes, "application/pdf")}
    )
    assert upload_res.status_code == 200
    doc_data = upload_res.json()
    assert doc_data["file_type"] == "PDF"
    assert doc_data["processing_status"] in ["PROCESSED", "REVIEW_REQUIRED"]
    assert doc_data["page_count"] == 2

    # Check document status endpoint
    doc_id = doc_data["id"]
    status_res = client.get(f"/api/documents/{doc_id}/status")
    assert status_res.status_code == 200
    status_data = status_res.json()
    assert status_data["topics_count"] >= 2

    # Verify extracted topics from PDF
    topics_res = client.get(f"/api/units/{unit_id}/topics")
    assert topics_res.status_code == 200
    topics = topics_res.json()
    assert len(topics) >= 2
    # First topic should be unlocked initially
    assert topics[0]["is_unlocked"] is True

    # Check topic content viewer
    content_res = client.get(f"/api/topics/{topics[0]['id']}/content")
    assert content_res.status_code == 200
    content = content_res.json()
    assert len(content["regions"]) > 0
    assert any(r["region_type"] in ["FULL_PAGE", "TEXT"] for r in content["regions"])

def test_pptx_processing(client, sample_pptx_bytes):
    s_res = client.post("/api/subjects", json={"name": "Computer Science"})
    u_res = client.post(f"/api/subjects/{s_res.json()['id']}/units", json={"title": "Unit 1: Computing Basics"})
    unit_id = u_res.json()["id"]

    # Upload PPTX
    upload_res = client.post(
        f"/api/units/{unit_id}/document",
        files={"file": ("slides.pptx", sample_pptx_bytes, "application/vnd.openxmlformats-officedocument.presentationml.presentation")}
    )
    assert upload_res.status_code == 200
    doc_data = upload_res.json()
    assert doc_data["file_type"] == "PPTX"
    assert doc_data["slide_count"] == 3

    # Check topics preserving slide order
    topics_res = client.get(f"/api/units/{unit_id}/topics")
    assert topics_res.status_code == 200
    topics = topics_res.json()
    assert len(topics) == 3
    # Verify exact slide titles preserved
    assert "Introduction to Computing" in topics[0]["name"]
    assert "Computer Architecture" in topics[1]["name"]
    assert "Operating Systems" in topics[2]["name"]
    # Check ordering
    assert topics[0]["order_index"] == 1
    assert topics[1]["order_index"] == 2
    assert topics[2]["order_index"] == 3

def test_docx_processing(client, sample_docx_bytes):
    s_res = client.post("/api/subjects", json={"name": "Operating Systems Course"})
    u_res = client.post(f"/api/subjects/{s_res.json()['id']}/units", json={"title": "Unit 1: OS Principles"})
    unit_id = u_res.json()["id"]

    # Upload DOCX
    upload_res = client.post(
        f"/api/units/{unit_id}/document",
        files={"file": ("notes.docx", sample_docx_bytes, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}
    )
    assert upload_res.status_code == 200
    doc_data = upload_res.json()
    assert doc_data["file_type"] == "DOCX"

    # Check topics preserving Heading 1 & 2 hierarchy
    topics_res = client.get(f"/api/units/{unit_id}/topics")
    assert topics_res.status_code == 200
    topics = topics_res.json()
    assert len(topics) >= 2
    assert "Operating System Fundamentals" in topics[0]["name"]
    assert len(topics[0]["subtopics"]) >= 1
    assert "Process Management" in topics[0]["subtopics"][0]["name"]

def test_semantic_topic_boundary_and_multi_slide_merging(client):
    """
    Test the fundamental Golden Rule:
    Slides/pages are source units; topics are semantic learning units.
    Verify:
    Slides 1-2 Planning
    Slides 3-4 Characteristics of Planning
    Slides 5-6 Planning Process
    Slide 7 Types of Plans
    -> Exactly 4 topics, NOT 7 topics!
    """
    from pptx import Presentation

    prs = Presentation()

    # Slide 1: Planning
    s1 = prs.slides.add_slide(prs.slide_layouts[1])
    s1.shapes.title.text = "Planning"
    s1.placeholders[1].text = "Planning is the foundational management function deciding what to do and how to do it."

    # Slide 2: Planning continuation (e.g. Planning Continued)
    s2 = prs.slides.add_slide(prs.slide_layouts[1])
    s2.shapes.title.text = "Planning (Continued)"
    s2.placeholders[1].text = "Why Planning is Important: It provides direction, reduces uncertainty and minimizes waste."

    # Slide 3: Characteristics of Planning
    s3 = prs.slides.add_slide(prs.slide_layouts[1])
    s3.shapes.title.text = "Characteristics of Planning"
    s3.placeholders[1].text = "Key features include goal orientation, pervasive presence, and futuristic outlook."

    # Slide 4: Characteristics continuation (e.g. Characteristics Part 2)
    s4 = prs.slides.add_slide(prs.slide_layouts[1])
    s4.shapes.title.text = "Characteristics of Planning (Part 2)"
    s4.placeholders[1].text = "Continuous process, decision-making exercise, and a mental exercise requiring imagination."

    # Slide 5: Planning Process
    s5 = prs.slides.add_slide(prs.slide_layouts[1])
    s5.shapes.title.text = "Planning Process"
    s5.placeholders[1].text = "Planning involves a systematic multi-stage workflow from setting goals to follow-up."

    # Slide 6: Steps in Planning Process (Process continuation)
    s6 = prs.slides.add_slide(prs.slide_layouts[1])
    s6.shapes.title.text = "Steps in Planning Process"
    s6.placeholders[1].text = "1. Setting objectives\n2. Developing premises\n3. Identifying alternatives\n4. Evaluating alternatives."

    # Slide 7: Types of Plans (New topic)
    s7 = prs.slides.add_slide(prs.slide_layouts[1])
    s7.shapes.title.text = "Types of Plans"
    s7.placeholders[1].text = "Plans can be classified as single-use plans and standing plans (policies, procedures, rules)."

    buf = io.BytesIO()
    prs.save(buf)
    pptx_bytes = buf.getvalue()

    # Create subject and unit
    s_res = client.post("/api/subjects", json={"name": "Principles of Management"})
    subj_id = s_res.json()["id"]
    u_res = client.post(f"/api/subjects/{subj_id}/units", json={"title": "Unit 2: Planning Functions"})
    unit_id = u_res.json()["id"]

    # Upload the 7-slide PPTX
    upload_res = client.post(
        f"/api/units/{unit_id}/document",
        files={"file": ("planning_lecture.pptx", pptx_bytes, "application/vnd.openxmlformats-officedocument.presentationml.presentation")}
    )
    assert upload_res.status_code == 200
    doc_data = upload_res.json()
    assert doc_data["slide_count"] == 7

    # Verify extracted topics: Exactly 4 topics, NOT 7!
    topics_res = client.get(f"/api/units/{unit_id}/topics")
    assert topics_res.status_code == 200
    topics = topics_res.json()

    assert len(topics) == 4, f"Expected exactly 4 topics from 7 slides, but got {len(topics)}: {[t['name'] for t in topics]}"

    # Verify Topic 1: Planning (Slides 1-2)
    assert "planning" in topics[0]["name"].lower()
    assert topics[0]["source_start"] == "1"
    assert topics[0]["source_end"] == "2"

    # Verify Topic 2: Characteristics of Planning (Slides 3-4)
    assert "characteristic" in topics[1]["name"].lower()
    assert topics[1]["source_start"] == "3"
    assert topics[1]["source_end"] == "4"

    # Verify Topic 3: Planning Process (Slides 5-6)
    assert "process" in topics[2]["name"].lower()
    assert topics[2]["source_start"] == "5"
    assert topics[2]["source_end"] == "6"

    # Verify Topic 4: Types of Plans (Slide 7)
    assert "types" in topics[3]["name"].lower()
    assert topics[3]["source_start"] == "7"
    assert topics[3]["source_end"] == "7"

