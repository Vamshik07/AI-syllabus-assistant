import pytest
from app.services.validation_service import validation_service
from app.services.structure_service import structure_service


def test_pptx_index_slide_is_not_created_as_topic():
    parsed_doc = {
        "slides": [
            {
                "slide_number": 1,
                "title": "UNIT-II Planning and Decision making",
                "text": (
                    "Title: UNIT-II Planning and Decision making\n"
                    "General Framework for Planning\n"
                    "Planning Process\n"
                    "Types of Plans\n"
                    "Management by Objectives\n"
                    "Production Planning and Control\n"
                ),
            },
            {
                "slide_number": 2,
                "title": "General Framework for Planning",
                "text": "Title: General Framework for Planning\nPlanning establishes organizational direction.",
            },
            {
                "slide_number": 3,
                "title": "Planning Process",
                "text": "Title: Planning Process\nPlanning follows a structured sequence.",
            },
        ],
        "full_text": (
            "General Framework for Planning\n"
            "Planning Process\n"
            "Types of Plans\n"
            "Management by Objectives\n"
            "Production Planning and Control"
        ),
    }

    topics, status = structure_service.extract_structure("PPTX", parsed_doc)

    assert status == "VERIFIED"
    assert parsed_doc["index"]["slide_number"] == 1
    assert parsed_doc["index"]["entries"][:3] == [
        "General Framework for Planning",
        "Planning Process",
        "Types of Plans",
    ]
    assert [topic["name"] for topic in topics] == [
        "General Framework for Planning",
        "Planning Process",
    ]
    assert topics[0]["source_start"] == "2"


def test_image_only_visual_headings_create_topics_without_slide_topics():
    parsed_doc = {
        "slides": [
            {"slide_number": 1, "title": "SDMA", "text": "Title: SDMA"},
            {
                "slide_number": 2,
                "title": "SDMA Architecture",
                "text": "",
                "visual_title": True,
                "elements": [{"type": "IMAGE", "visual_content": "Architecture labels"}],
            },
            {
                "slide_number": 3,
                "title": "SDMA Working",
                "text": "",
                "visual_title": True,
                "elements": [{"type": "IMAGE", "visual_content": "Working stages"}],
            },
        ],
        "full_text": "Title: SDMA",
    }

    topics, status = structure_service.extract_structure("PPTX", parsed_doc)

    assert status == "VERIFIED"
    assert [topic["name"] for topic in topics] == [
        "SDMA",
        "SDMA Architecture",
        "SDMA Working",
    ]
    assert [(topic["source_start"], topic["source_end"]) for topic in topics] == [
        ("1", "1"),
        ("2", "2"),
        ("3", "3"),
    ]


def test_multi_slide_topic_keeps_complete_source_content(client):
    from pptx import Presentation
    import io

    prs = Presentation()
    for title, body in [
        ("Planning Process", "Planning establishes objectives and evaluates alternatives."),
        ("Planning Process", "The selected plan is implemented and reviewed against objectives."),
        ("Types of Plans", "Strategic plans define long-term direction."),
    ]:
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = title
        slide.placeholders[1].text = body

    buf = io.BytesIO()
    prs.save(buf)
    subject = client.post("/api/subjects", json={"name": "Multi Slide Subject"}).json()
    unit = client.post(f"/api/subjects/{subject['id']}/units", json={"title": "Planning Unit"}).json()
    upload = client.post(
        f"/api/units/{unit['id']}/document",
        files={"file": ("planning.pptx", buf.getvalue(), "application/vnd.openxmlformats-officedocument.presentationml.presentation")},
    )
    assert upload.status_code == 200

    topics = client.get(f"/api/units/{unit['id']}/topics").json()
    assert [topic["name"] for topic in topics] == ["Planning Process", "Types of Plans"]
    assert topics[0]["source_start"] == "1"
    assert topics[0]["source_end"] == "2"

    content = client.get(f"/api/topics/{topics[0]['id']}/content").json()
    sections = content["learning_data"]["source_sections"]
    source_text = " ".join(section["text"] for section in sections)
    assert "evaluates alternatives" in source_text
    assert "implemented and reviewed" in source_text
    assert all(section["slide_number"] in (1, 2) for section in sections)

def test_source_fidelity_topic_names_and_ordering(client, sample_pptx_bytes, sample_docx_bytes):
    # Test PPTX source grounding
    s1 = client.post("/api/subjects", json={"name": "Fidelity Check Subj"}).json()
    u1 = client.post(f"/api/subjects/{s1['id']}/units", json={"title": "Fidelity Unit 1"}).json()
    
    upload_res = client.post(
        f"/api/units/{u1['id']}/document",
        files={"file": ("slides.pptx", sample_pptx_bytes, "application/vnd.openxmlformats-officedocument.presentationml.presentation")}
    )
    doc_id = upload_res.json()["id"]

    topics = client.get(f"/api/units/{u1['id']}/topics").json()
    
    # Requirement: Topic order matches source order
    order_indices = [t["order_index"] for t in topics]
    assert order_indices == sorted(order_indices), "Topic order must strictly increase"

    # Requirement: Names must be authentic
    expected_pptx_titles = ["Introduction to Computing", "Computer Architecture", "Operating Systems"]
    for idx, expected in enumerate(expected_pptx_titles):
        assert expected in topics[idx]["name"]

    # Test DOCX source grounding
    u2 = client.post(f"/api/subjects/{s1['id']}/units", json={"title": "Fidelity Unit 2"}).json()
    client.post(
        f"/api/units/{u2['id']}/document",
        files={"file": ("notes.docx", sample_docx_bytes, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}
    )
    docx_topics = client.get(f"/api/units/{u2['id']}/topics").json()

    # Verify topics and subtopics exist verbatim in source
    assert "Operating System Fundamentals" in docx_topics[0]["name"]
    assert docx_topics[0]["order_index"] == 1
    assert "Memory Management" in docx_topics[1]["name"]
    assert docx_topics[1]["order_index"] == 2

    # Verify subtopics
    assert len(docx_topics[0]["subtopics"]) > 0
    assert "Process Management" in docx_topics[0]["subtopics"][0]["name"]
    assert docx_topics[0]["subtopics"][0]["order_index"] == 1

def test_quiz_source_isolation_and_grounding(client, sample_pptx_bytes):
    s = client.post("/api/subjects", json={"name": "Quiz Isolation Subj"}).json()
    u = client.post(f"/api/subjects/{s['id']}/units", json={"title": "Quiz Isolation Unit"}).json()
    client.post(
        f"/api/units/{u['id']}/document",
        files={"file": ("slides.pptx", sample_pptx_bytes, "application/vnd.openxmlformats-officedocument.presentationml.presentation")}
    )
    topics = client.get(f"/api/units/{u['id']}/topics").json()
    
    # Topic 1: Introduction to Computing
    # Topic 2: Computer Architecture
    t1_id = topics[0]["id"]
    t2_id = topics[1]["id"]

    # Generate quiz for Topic 1
    q1_res = client.post(f"/api/topics/{t1_id}/quiz/generate")
    assert q1_res.status_code == 200
    q1_data = q1_res.json()
    assert len(q1_data["questions"]) == 5

    # Get Topic 1 source content
    t1_content = client.get(f"/api/topics/{t1_id}/content").json()
    t1_text = " ".join([r.get("content_text") or "" for r in t1_content["regions"]]).lower()

    # Topic 2 source content
    # Note: Topic 2 is initially locked, but we can verify that Topic 1's questions do NOT mention Topic 2's specific concepts
    for q in q1_data["questions"]:
        assert q["source_reference"] is not None
        # Verify question is grounded in Topic 1 source
        q_ref = q["source_reference"].lower()
        # Words in reference must belong to Topic 1
        ref_words = [w for w in q_ref.split() if len(w) > 4]
        if ref_words:
            assert any(w in t1_text for w in ref_words), f"Question source reference '{q_ref}' not in Topic 1 text"
        # Von Neumann is exclusively in Topic 2 (Computer Architecture), so it must NOT be the source of Topic 1 questions
        assert "von neumann" not in q["question"].lower()

def test_presenter_boilerplate_filtered_and_no_mindmap_flowchart(client):
    """
    Verifies that presenter/faculty boilerplate (e.g. 'Dr. B. Rebecca, Associate Professor')
    is NEVER extracted as a core concept or definition, and that mind_map and flowchart
    are completely removed.
    """
    from pptx import Presentation
    import io
    from app.services.learning_service import learning_service
    from app.models.models import Topic

    prs = Presentation()
    # Slide 1: Topic with presenter info
    s1 = prs.slides.add_slide(prs.slide_layouts[1])
    s1.shapes.title.text = "Space Division Multiple Access"
    s1.placeholders[1].text = "Dr. B. Rebecca, Associate Professor\nDepartment of Computer Science and Engineering\nCollege of Engineering"

    # Slide 2: Educational features and advantages
    s2 = prs.slides.add_slide(prs.slide_layouts[1])
    s2.shapes.title.text = "SDMA Features"
    s2.placeholders[1].text = (
        "Space Division Multiple Access is defined as an advanced channel allocation scheme utilizing spatial separation.\n"
        "Advantage: Increases wireless capacity significantly without additional spectrum.\n"
        "Key characteristic: Utilizes smart directional antennas to form radiation beams."
    )

    buf = io.BytesIO()
    prs.save(buf)
    pptx_bytes = buf.getvalue()

    s = client.post("/api/subjects", json={"name": "Mobile Computing Course"}).json()
    u = client.post(f"/api/subjects/{s['id']}/units", json={"title": "Unit 2: Multiple Access"}).json()

    upload_res = client.post(
        f"/api/units/{u['id']}/document",
        files={"file": ("sdma_lecture.pptx", pptx_bytes, "application/vnd.openxmlformats-officedocument.presentationml.presentation")}
    )
    assert upload_res.status_code == 200

    topics = client.get(f"/api/units/{u['id']}/topics").json()
    assert len(topics) >= 1
    t = topics[0]

    # Get topic content / learning data
    content = client.get(f"/api/topics/{t['id']}/content").json()
    learning_data = content.get("learning_data", {})

    # 1. Verify Mind Map and Flowchart are completely absent
    assert "mind_map" not in learning_data, "mind_map must not be present in learning_data"
    assert "flowchart" not in learning_data, "flowchart must not be present in learning_data"

    # 2. Verify presenter name and affiliation are NEVER extracted as concepts or definitions
    serialized_data = str(learning_data).lower()
    assert "rebecca" not in serialized_data, "Presenter name must never be extracted into learning data"
    assert "associate professor" not in serialized_data, "Faculty rank must never be extracted into learning data"
    assert "department of" not in serialized_data, "Department boilerplate must never be extracted into learning data"

    # 3. Verify real definitions and advantages ARE extracted
    defs = learning_data.get("definitions", [])
    advs = learning_data.get("advantages", [])
    assert len(defs) > 0 or len(learning_data.get("core_concepts", [])) > 0
    assert any("wireless capacity" in str(a).lower() for a in advs) or any("capacity" in str(c).lower() for c in learning_data.get("core_concepts", []))


def test_pptx_embedded_image_text_is_not_learning_source(client):
    """Embedded slide images remain visuals and are not OCR-fed into lesson text."""
    from pptx import Presentation
    from pptx.util import Inches
    from PIL import Image, ImageDraw
    import io

    image = Image.new("RGB", (800, 300), "white")
    ImageDraw.Draw(image).text((40, 120), "PIXEL ONLY DIAGRAM LABEL", fill="black")
    image_buffer = io.BytesIO()
    image.save(image_buffer, format="PNG")

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "SDMA Overview"
    slide.placeholders[1].text = "SDMA separates users spatially to improve wireless capacity."
    slide.shapes.add_picture(io.BytesIO(image_buffer.getvalue()), Inches(1), Inches(3), width=Inches(7))

    pptx_buffer = io.BytesIO()
    prs.save(pptx_buffer)

    subject = client.post("/api/subjects", json={"name": "Mobile Computing"}).json()
    unit = client.post(
        f"/api/subjects/{subject['id']}/units",
        json={"title": "Unit 1"},
    ).json()
    upload = client.post(
        f"/api/units/{unit['id']}/document",
        files={
            "file": (
                "sdma-image.pptx",
                pptx_buffer.getvalue(),
                "application/vnd.openxmlformats-officedocument.presentationml.presentation",
            )
        },
    )
    assert upload.status_code == 200

    topics = client.get(f"/api/units/{unit['id']}/topics").json()
    content = client.get(f"/api/topics/{topics[0]['id']}/content").json()
    learning_data = content["learning_data"]
    # Image pixels are never used to name a topic or replace the slide's
    # authoritative native text. If visual analysis is available, its output
    # may appear only as supplemental IMAGE evidence.
    assert topics[0]["name"] == "SDMA Overview"
    slide_regions = [r for r in content["regions"] if r["region_type"] == "SLIDE"]
    assert "pixel only diagram label" not in " ".join(
        r.get("content_text") or "" for r in slide_regions
    ).lower()
    assert any(region["region_type"] == "IMAGE" for region in content["regions"])
    assert any(visual["type"] == "IMAGE" for visual in learning_data["visuals"])
